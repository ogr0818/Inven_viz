#!/usr/bin/env python3
"""將 .xlsx 月份工作表名稱轉為民國年月，並輸出指定欄位。

安裝：python3 -m pip install openpyxl
使用：python3 rename_monthly_xlsx.py "2026年每月報表.xlsx"
固定輸出至執行目錄下的 ./月報表.xlsx；重複執行會覆寫此輸出檔。
每張工作表只輸出 drug_id、實發量、庫存量，依此順序排列。
各工作表內相同 drug_id 合併，實發量、庫存量分別加總。
數量空白視為 0；非數字數量或缺少 drug_id 的資料列會回報錯誤。
標題列預設為第 1 列，可使用 --header-row 指定。

不符合「西元年年月份月」的工作表名稱保持原樣。
輸出為資料值；公式欄位使用 Excel 最後儲存的計算結果。
"""

import argparse
import math
from pathlib import Path
import re

from openpyxl import Workbook, load_workbook


HEADER_MAPPING = {"藥品代碼": "drug_id", "藥品名稱": "drug_name"}
OUTPUT_COLUMNS = ["drug_id", "實發量", "庫存量"]
OUTPUT_PATH = Path("./月報表.xlsx")


def month_sheet_name(name: str) -> str:
    """例如：2026年1月 → 11501、2026年12月 → 11512。"""
    match = re.fullmatch(r"\s*([0-9]{4})\s*年\s*([0-9]{1,2})\s*月\s*", name)
    if match is None:
        return name
    year, month = map(int, match.groups())
    if year <= 1911 or not 1 <= month <= 12:
        raise ValueError(f"工作表年月無效：{name}")
    return f"{year - 1911:03d}{month:02d}"


def is_blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def aggregate_rows(sheet, indices, header_row):
    """按代碼首次出現順序，累積同一工作表內的兩欄數量。"""
    totals = {}
    for row_number, row in enumerate(
        sheet.iter_rows(min_row=header_row + 1, values_only=True),
        start=header_row + 1,
    ):
        if all(is_blank(value) for value in row):
            continue
        drug_id, *quantities = [row[index] for index in indices]
        if is_blank(drug_id):
            raise ValueError(f"工作表 {sheet.title} 第 {row_number} 列欄位 drug_id 缺少值。")
        for index, value in enumerate(quantities):
            if is_blank(value):
                quantities[index] = 0
            elif (isinstance(value, bool) or not isinstance(value, (int, float))
                  or (isinstance(value, float) and not math.isfinite(value))):
                column = OUTPUT_COLUMNS[index + 1]
                raise ValueError(
                    f"工作表 {sheet.title} 第 {row_number} 列欄位 {column} "
                    f"必須為數字，收到：{value!r}"
                )
        sums = totals.setdefault(drug_id, [0, 0])
        sums[0] += quantities[0]
        sums[1] += quantities[1]
    return totals


def convert_xlsx(input_path, header_row=1):
    """處理所有工作表並另存檔案；回傳輸出 Path。"""
    source = Path(input_path).expanduser().resolve()
    target = OUTPUT_PATH.resolve()
    if source.suffix.lower() != ".xlsx" or target.suffix.lower() != ".xlsx":
        raise ValueError("輸入與輸出檔案皆須為 .xlsx。")
    if not source.is_file():
        raise FileNotFoundError(f"找不到輸入檔案：{source}")
    if header_row < 1:
        raise ValueError("標題列必須大於或等於 1。")
    if source == target:
        raise ValueError("輸出路徑必須與原始檔案不同。")
    workbook = load_workbook(source, read_only=True, data_only=True)
    result = Workbook(write_only=True)
    try:
        names = {sheet.title: month_sheet_name(sheet.title) for sheet in workbook.worksheets}
        # 先檢查名稱衝突，避免 Excel 自動添加尾碼。
        if len({name.casefold() for name in names.values()}) != len(names):
            raise ValueError("轉換後的工作表名稱重複，請先調整原始名稱。")

        # 先驗證所有工作表，缺欄位或重複欄位時不覆寫既有輸出。
        selections = []
        for sheet in workbook.worksheets:
            headers = next(sheet.iter_rows(min_row=header_row, max_row=header_row,
                                           values_only=True), ())
            normalized = [HEADER_MAPPING.get(value.strip(), value.strip())
                          if isinstance(value, str) else value for value in headers]
            for column in OUTPUT_COLUMNS:
                if normalized.count(column) != 1:
                    raise ValueError(f"工作表 {sheet.title} 的欄位 {column} 缺少或重複。")
            selections.append((sheet, [normalized.index(column) for column in OUTPUT_COLUMNS]))

        # 所有資料驗證成功後才開始輸出，避免錯誤資料覆寫既有檔案。
        aggregated = [(sheet, aggregate_rows(sheet, indices, header_row))
                      for sheet, indices in selections]
        for sheet, totals in aggregated:
            output_sheet = result.create_sheet(names[sheet.title])
            output_sheet.append(OUTPUT_COLUMNS)
            for drug_id, quantities in totals.items():
                output_sheet.append([drug_id, *quantities])
            print(f"{sheet.title} → {output_sheet.title}；輸出 {len(totals)} 筆資料")

        result.save(target)
    finally:
        workbook.close()
        result.close()
    return target


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", help="來源 .xlsx 路徑")
    parser.add_argument("--header-row", type=int, default=1, help="標題列，預設 1")
    args = parser.parse_args()
    try:
        output = convert_xlsx(args.input, args.header_row)
    except (OSError, ValueError) as error:
        parser.exit(1, f"錯誤：{error}\n")
    print(f"已儲存：{output}")


if __name__ == "__main__":
    main()
