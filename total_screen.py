"""讀取 OUT、UD 的所有工作表，依年月及藥品代碼加總原始數量。"""

import math
from pathlib import Path
import sys
import tempfile
from zipfile import BadZipFile

import pandas as pd
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill

from quantity_rules import filter_quantity_range


BASE_DIR = Path(__file__).resolve().parent
INPUT_NAMES = ("out_screen.xlsx", "ud_screen.xlsx")
OUTPUT_NAME = "total_screen.xlsx"
COLUMNS = ["drug_id", "total_qty"]
OUTPUT_COLUMNS = [*COLUMNS, "inv_year", "inv_month"]
GROUP_COLUMNS = ["inv_year", "inv_month", "drug_id"]
MAX_DATA_ROWS = 1_048_575


def inventory_period(sheet: str) -> tuple[str, str]:
    """從工作表名稱取得三位年份及兩位月份文字。"""
    year, month = sheet[:3], sheet[3:]
    if (
        len(year) != 3
        or not year.isascii()
        or not year.isdigit()
        or len(month) not in (1, 2)
        or not month.isascii()
        or not month.isdigit()
        or not 1 <= int(month) <= 12
    ):
        raise ValueError(f"工作表 {sheet}：名稱必須是三位年份加月份（1 至 12），例如 11501")
    return year, month.zfill(2)


def row_numbers(mask: pd.Series) -> str:
    """回報最多五個 Excel 列號，標題列為第 1 列。"""
    indices = mask[mask].index
    rows = ", ".join(str(int(index) + 2) for index in indices[:5])
    return rows + (" 等" if len(indices) > 5 else "")


def read_sheet(book: pd.ExcelFile, path: Path, sheet: str) -> pd.DataFrame:
    """只讀必要欄位；代碼保留原文，數量必須為有限數值。"""
    context = f"{path.name}／工作表 {sheet}"
    try:
        frame = pd.read_excel(
            book,
            sheet_name=sheet,
            usecols=lambda column: column in COLUMNS,
            dtype="string",
            keep_default_na=False,
        )
    except (OSError, ValueError) as exc:
        raise ValueError(f"{context}：無法讀取：{exc}") from exc
    missing = [column for column in COLUMNS if column not in frame.columns]
    if missing:
        raise ValueError(f"{context}：缺少必要欄位：{', '.join(missing)}")

    frame = frame[COLUMNS].copy()
    frame = filter_quantity_range(frame, context)
    invalid_codes = frame["drug_id"].str.contains(ILLEGAL_CHARACTERS_RE, na=False)
    if invalid_codes.any():
        raise ValueError(f"{context}：第 {row_numbers(invalid_codes)} 列 drug_id 含 Excel 不允許的控制字元")
    blank = frame["drug_id"].isna() | frame["drug_id"].str.strip().eq("")
    if blank.any():
        raise ValueError(f"{context}：第 {row_numbers(blank)} 列 drug_id 空白")
    return frame


def total_to_excel(base_dir: Path = BASE_DIR) -> pd.DataFrame:
    """逐表彙總，成功寫入暫存檔後才取代固定輸出。"""
    sources = [base_dir / name for name in INPUT_NAMES]
    for path in sources:
        if not path.is_file():
            raise ValueError(f"找不到輸入檔案：{path}")

    subtotals = []
    for path in sources:
        try:
            with pd.ExcelFile(path, engine="openpyxl") as book:
                for sheet in book.sheet_names:
                    year, month = inventory_period(sheet)
                    frame = read_sheet(book, path, sheet)
                    if not frame.empty:
                        subtotal = frame.groupby("drug_id", as_index=False, sort=False)["total_qty"].sum()
                        subtotal["inv_year"] = pd.Series(year, index=subtotal.index, dtype="string")
                        subtotal["inv_month"] = pd.Series(month, index=subtotal.index, dtype="string")
                        subtotals.append(subtotal)
                    ignored = frame.attrs["ignored_quantity_rows"]
                    print(f"{path.name}／{sheet}：保留 {len(frame):,} 筆；數量超出範圍忽略 {ignored:,} 筆", flush=True)
        except (OSError, ValueError, BadZipFile) as exc:
            raise ValueError(f"{path.name}：{exc}") from exc

    if subtotals:
        result = (
            pd.concat(subtotals, ignore_index=True)
            .groupby(GROUP_COLUMNS, as_index=False, sort=True)["total_qty"]
            .sum()
        )[OUTPUT_COLUMNS]
        if not result["total_qty"].map(math.isfinite).all():
            raise ValueError("加總後 total_qty 超出有效數值範圍")
    else:
        result = pd.DataFrame(columns=OUTPUT_COLUMNS)
    for column in ("drug_id", "inv_year", "inv_month"):
        result[column] = result[column].astype("string")
    if len(result) > MAX_DATA_ROWS:
        raise ValueError(f"彙總結果超過 Excel 單張工作表 {MAX_DATA_ROWS:,} 筆資料的上限")

    output = base_dir / OUTPUT_NAME
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=base_dir, suffix=".xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
            # 預先建立工作表，避免寫入失敗時關閉 writer 又產生無工作表錯誤。
            writer.book.create_sheet("total_screen")
            result.to_excel(writer, sheet_name="total_screen", index=False)
            worksheet = writer.sheets["total_screen"]
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            worksheet.sheet_view.showGridLines = False
            worksheet.column_dimensions["A"].width = max(
                14, max((len(str(value)) + 2 for value in result["drug_id"]), default=0)
            )
            worksheet.column_dimensions["B"].width = 22
            worksheet.column_dimensions["C"].width = 14
            worksheet.column_dimensions["D"].width = 14
            for cell in worksheet[1]:
                cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="334155")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for code, quantity, year, month in worksheet.iter_rows(min_row=2, max_row=len(result) + 1):
                # 文字代碼（包含以 = 開頭的代碼）不可被寫成公式。
                for cell in (code, year, month):
                    cell.data_type = "s"
                    cell.number_format = "@"
                    cell.font = Font(name="Arial", size=10)
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                quantity.font = Font(name="Arial", size=10)
                quantity.alignment = Alignment(horizontal="right", vertical="center")
        temporary.replace(output)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    print(f"完成：{output.resolve()}（共 {len(result):,} 筆年月藥品彙總）", flush=True)
    return result


def main() -> int:
    try:
        total_to_excel()
    except (OSError, ValueError, BadZipFile) as exc:
        print(f"錯誤：{exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
