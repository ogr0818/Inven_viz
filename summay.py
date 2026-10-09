#!/usr/bin/env python3
"""以月報為主，合併基本資料與同年月耗量，輸出中文欄位月明細。

固定讀取本腳本所在資料夾的 base.xlsx、月報表.xlsx、total_screen.xlsx，
並將結果寫入同一資料夾的 summay.xlsx。在 VS Code 直接執行本檔即可：
    python summay.py

讀取路徑以本腳本位置為準，不受 VS Code 工作目錄影響。保留月報工作表名稱、
工作表順序及各表列順序；實發量小於 0 的整列排除，缺當月耗量補 0。
重複鍵停止處理，完整回報內容；所有驗證與寫入成功後才取代舊輸出。
"""

import math
from pathlib import Path
import re
import sys
import tempfile
import unicodedata
from zipfile import BadZipFile

import pandas as pd
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill


BASE_DIR = Path(__file__).resolve().parent
INPUT_NAMES = ("base.xlsx", "月報表.xlsx", "total_screen.xlsx")
OUTPUT_NAME = "summay.xlsx"
BASE_COLUMNS = ["drug_id", "drug_name", "drug_type"]
MONTHLY_COLUMNS = ["drug_id", "實發量", "庫存量"]
TOTAL_COLUMNS = ["drug_id", "total_qty", "inv_year", "inv_month"]
KEY_COLUMNS = ["drug_id", "inv_year", "inv_month"]
OUTPUT_NAMES = {
    "drug_id": "藥品代碼",
    "drug_name": "藥品名稱",
    "實發量": "實發量",
    "total_qty": "住院耗用",
    "庫存量": "庫存量",
    "inv_year": "撥補年份",
    "inv_month": "撥補月份",
    "drug_type": "劑型",
}
SOURCE_COLUMNS = ["來源檔案", "來源工作表", "Excel 列號"]
MAX_DATA_ROWS = 1_048_575


def inventory_period(sheet: str) -> tuple[str, str]:
    """解析民國年月；輸出保留原頁籤名，月份以兩碼文字儲存。"""
    if not re.fullmatch(r"[0-9]{3}[0-9]{1,2}", sheet):
        raise ValueError(f"月報工作表 {sheet}：名稱須為民國年份加月份，例如 11501")
    year, month = sheet[:3], sheet[3:]
    if not 1 <= int(month) <= 12:
        raise ValueError(f"月報工作表 {sheet}：月份須介於 1 至 12")
    return year, month.zfill(2)


def read_frame(book: pd.ExcelFile, path: Path, sheet: str, columns: list[str]) -> pd.DataFrame:
    """檢查原始標題列，保留文字值及原 Excel 列號。"""
    headers = list(next(book.book[sheet].iter_rows(min_row=1, max_row=1, values_only=True), ()))
    invalid = [column for column in columns if headers.count(column) != 1]
    if invalid:
        raise ValueError(f"{path.name}／{sheet}：必要欄位缺少或重複：{', '.join(invalid)}")
    frame = pd.read_excel(book, sheet_name=sheet, dtype="string", keep_default_na=False)[columns].copy()
    frame["來源檔案"] = path.name
    frame["來源工作表"] = sheet
    frame["Excel 列號"] = frame.index + 2
    return frame


def report_invalid(frame: pd.DataFrame, mask: pd.Series, message: str, columns: list[str]) -> None:
    if mask.any():
        details = frame.loc[mask, SOURCE_COLUMNS + columns].to_string(index=False, max_rows=None, max_cols=None)
        raise ValueError(f"{message}\n{details}")


def validate_text(frame: pd.DataFrame, columns: list[str]) -> None:
    for column in columns:
        values = frame[column]
        invalid = values.isna() | values.str.strip().eq("") | values.str.contains(ILLEGAL_CHARACTERS_RE, na=False)
        report_invalid(frame, invalid, f"{column} 空白或含不允許的控制字元：", columns)
    report_invalid(
        frame, ~frame["drug_id"].str.fullmatch(r"[A-Z0-9]{6}", na=False),
        "drug_id 須為六碼大寫英數文字：", columns,
    )


def validate_numbers(frame: pd.DataFrame, columns: list[str]) -> None:
    """只驗證有限數值，不套用上游交易數量界線，也不變更正負號。"""
    for column in columns:
        numbers = pd.to_numeric(frame[column], errors="coerce")
        valid = numbers.notna() & numbers.map(lambda value: math.isfinite(value) if pd.notna(value) else False)
        report_invalid(frame, ~valid, f"{column} 須為有效有限數值：", columns)
        frame[column] = numbers


def validate_unique(frame: pd.DataFrame, keys: list[str], label: str, raw: pd.DataFrame | None = None) -> None:
    duplicates = frame.duplicated(keys, keep=False)
    if duplicates.any():
        groups = len(frame.loc[duplicates, keys].drop_duplicates())
        rows = raw.loc[duplicates] if raw is not None else frame.loc[duplicates]
        fields = SOURCE_COLUMNS + [column for column in rows.columns if column not in SOURCE_COLUMNS]
        details = rows[fields].to_string(index=False, max_rows=None, max_cols=None)
        raise ValueError(
            f"{label} 發現重複鍵（{'＋'.join(keys)}）：{groups} 組、{len(rows)} 列。"
            f"停止合併；以下為所有重複資料列：\n{details}"
        )


def normalize_periods(frame: pd.DataFrame) -> None:
    report_invalid(frame, ~frame["inv_year"].str.fullmatch(r"[0-9]{3}", na=False),
                   "inv_year 須為三碼民國年份文字：", ["inv_year", "inv_month"])
    months = frame["inv_month"]
    numeric = pd.to_numeric(months, errors="coerce")
    valid = months.str.fullmatch(r"[0-9]{1,2}", na=False) & numeric.between(1, 12)
    report_invalid(frame, ~valid, "inv_month 須為 1 至 12：", ["inv_year", "inv_month"])
    frame["inv_month"] = months.str.zfill(2)


def display_width(text: str) -> int:
    return sum(2 if unicodedata.east_asian_width(char) in "WF" else 1 for char in text)


def format_sheet(worksheet) -> None:
    worksheet.freeze_panes = "C2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False
    for letter, width in zip("ABCDEFGH", [16, 58, 16, 16, 16, 16, 16, 12]):
        worksheet.column_dimensions[letter].width = width
    worksheet.row_dimensions[1].height = 28
    for cell in worksheet[1]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="334155")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(horizontal="right", vertical="center")
        for cell in row[2:5]:
            cell.number_format = "#,##0" if float(cell.value).is_integer() else "#,##0.##########"
        for index in (0, 1, 5, 6, 7):
            cell = row[index]
            # 名稱即使以 = 開頭，也必須保持文字；年月不遺失前導零。
            cell.data_type = "s"
            cell.number_format = "@"
            cell.alignment = Alignment(horizontal="left" if index < 2 else "center", vertical="center", wrap_text=index == 1)
        name = str(row[1].value or "")
        lines = sum(max(1, math.ceil(display_width(line) / 56)) for line in name.split("\n"))
        worksheet.row_dimensions[row[0].row].height = max(24, lines * 15 + 8)


def merge_to_excel() -> dict[str, dict[str, int]]:
    """驗證完整輸入、合併每個品項月份，再以暫存檔取代輸出。"""
    base_path, monthly_path, total_path = (BASE_DIR / name for name in INPUT_NAMES)
    output_path = BASE_DIR / OUTPUT_NAME
    base_sheet = "base"
    sources = (base_path, monthly_path, total_path)
    if output_path.resolve() in {path.resolve() for path in sources}:
        raise ValueError("輸出路徑不可與任何輸入檔案相同")
    for path in sources:
        if not path.is_file():
            raise ValueError(f"找不到輸入檔案：{path}")

    with pd.ExcelFile(base_path, engine="openpyxl") as book:
        if base_sheet not in book.sheet_names:
            raise ValueError(f"{base_path.name}：找不到基本資料工作表 {base_sheet}")
        base = read_frame(book, base_path, base_sheet, BASE_COLUMNS)
    validate_unique(base, ["drug_id"], "基本資料")
    validate_text(base, BASE_COLUMNS)

    total_frames = []
    with pd.ExcelFile(total_path, engine="openpyxl") as book:
        for sheet in book.sheet_names:
            total_frames.append(read_frame(book, total_path, sheet, TOTAL_COLUMNS))
    total_raw = pd.concat(total_frames, ignore_index=True)
    total = total_raw.copy()
    normalize_periods(total)
    # 先完整回報重複列，再驗證數量，避免其中一筆壞值遮蔽重複內容。
    validate_unique(total, KEY_COLUMNS, "耗量表", total_raw)
    validate_text(total, ["drug_id"])
    validate_numbers(total, ["total_qty"])

    months = {}
    with pd.ExcelFile(monthly_path, engine="openpyxl") as book:
        for sheet in book.sheet_names:
            year, month = inventory_period(sheet)
            frame = read_frame(book, monthly_path, sheet, MONTHLY_COLUMNS)
            frame["inv_year"], frame["inv_month"] = year, month
            months[sheet] = frame
    validate_unique(pd.concat(months.values(), ignore_index=True), KEY_COLUMNS, "月報")

    outputs, counts = {}, {}
    for sheet, frame in months.items():
        validate_text(frame, ["drug_id"])
        validate_numbers(frame, ["實發量", "庫存量"])
        missing_base = ~frame["drug_id"].isin(base["drug_id"])
        report_invalid(frame, missing_base, "月報代碼在基本資料中不存在：", ["drug_id"])
        negative = frame["實發量"].lt(0)
        eligible = frame.loc[~negative].copy()
        merged = eligible.merge(base[BASE_COLUMNS], on="drug_id", how="left", sort=False, validate="many_to_one")
        merged = merged.merge(total[TOTAL_COLUMNS], on=KEY_COLUMNS, how="left", sort=False,
                              validate="one_to_one", indicator=True)
        missing_total = merged["_merge"].eq("left_only")
        merged.loc[missing_total, "total_qty"] = 0
        output = merged[list(OUTPUT_NAMES)].rename(columns=OUTPUT_NAMES)
        if len(output) > MAX_DATA_ROWS:
            raise ValueError(f"{sheet}：輸出超過 Excel 每張工作表的資料列上限")
        outputs[sheet] = output
        counts[sheet] = {"input": len(frame), "excluded": int(negative.sum()),
                         "filled_zero": int(missing_total.sum()), "output": len(output)}

    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=output_path.parent, prefix=".summay-", suffix=".xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
            writer.book.create_sheet(next(iter(outputs)))
            for sheet, frame in outputs.items():
                frame.to_excel(writer, sheet_name=sheet, index=False)
                format_sheet(writer.sheets[sheet])
        temporary.replace(output_path)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)

    for sheet, count in counts.items():
        print(f"{sheet}：月報 {count['input']:,} 筆；負實發量排除 {count['excluded']:,} 筆；"
              f"缺耗量補零 {count['filled_zero']:,} 筆；輸出 {count['output']:,} 筆")
    print(f"完成：{output_path.resolve()}（{len(outputs)} 張工作表，共 {sum(c['output'] for c in counts.values()):,} 筆月明細）")
    return counts


def main() -> int:
    try:
        merge_to_excel()
    except (OSError, ValueError, BadZipFile) as exc:
        print(f"錯誤：{exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
