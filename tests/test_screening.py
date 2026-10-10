"""數量界線、文字驗證及輸出保護的回歸測試。"""

import contextlib
import io
from pathlib import Path
import tempfile
import unittest

import pandas as pd
from openpyxl import load_workbook

import out_screen
from quantity_rules import filter_quantity_range
import total_screen
import ud_screen


class ScreeningTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.base = Path(self.temporary.name)

    def csv(self, module, rows, name="case.csv", encoding="utf-8-sig"):
        path = self.base / name
        pd.DataFrame(rows).reindex(columns=module.COLUMNS).to_csv(
            path, index=False, encoding=encoding
        )
        return path

    def row(self, quantity="1", code="001", name="drug", transaction="3", outlet="5630"):
        return dict(drug_id=code, drug_name=name, total_qty=quantity,
                    phtxid=transaction, phoutid=outlet)

    def test_exact_boundaries_and_large_numbers(self):
        values = ["-1000000", "1000000", "0", "1.25", "-1000001", "1000001",
                  "-9223372036854775808", "18446744073709551615", "1e400",
                  "1000000.00000000001", "-1000000.00000000001"]
        result = filter_quantity_range(pd.DataFrame({"total_qty": values}), "source")
        self.assertEqual(result.total_qty.tolist(), [-1000000, 1000000, 0, 1.25])
        self.assertEqual(result.attrs["ignored_quantity_rows"], 7)

    def test_invalid_quantities_fail(self):
        for value in ["", " ", "abc", "NaN", "Infinity", "-inf", "1_000"]:
            with self.subTest(value=value), self.assertRaisesRegex(ValueError, "第 2 .*total_qty"):
                filter_quantity_range(pd.DataFrame({"total_qty": [value]}), "source")

    def test_filtered_quantity_dtype_is_signed(self):
        result = filter_quantity_range(
            pd.DataFrame({"total_qty": ["1", "18446744073709551615"]}), "source"
        )
        self.assertEqual((-result.total_qty).tolist(), [-1])

    def test_csv_range_filter_and_signs(self):
        for module in [out_screen, ud_screen]:
            for encoding in ["utf-8-sig", "utf-8", "utf-16"]:
                with self.subTest(module=module.__name__, encoding=encoding):
                    rows = [self.row("-1000000", transaction="2"),
                            self.row("-1.25", code="002", transaction="3"),
                            self.row("1000001", code="", name="bad\x01", transaction="bad"),
                            self.row("invalid", outlet="9999")]
                    frame = module.read_screened_csv(self.csv(module, rows, encoding=encoding))
                    self.assertEqual(frame.drug_id.tolist(), ["001", "002"])
                    self.assertEqual(frame.total_qty.tolist(),
                                     [-1000000, 1.25] if module is ud_screen else [-1000000, -1.25])
                    self.assertEqual(frame.attrs["ignored_quantity_rows"], 1)

    def test_empty_and_all_out_of_range_export(self):
        for module in [out_screen, ud_screen]:
            for rows in [[], [self.row("1e400")]]:
                with self.subTest(module=module.__name__, rows=rows):
                    self.csv(module, rows, name="11501.csv")
                    output = self.base / "output.xlsx"
                    with contextlib.redirect_stdout(io.StringIO()):
                        counts = module.screen_to_excel(self.base, output)
                    self.assertEqual(counts, {"11501.csv": 0})
                    book = load_workbook(output, read_only=True)
                    try:
                        columns = module.OUTPUT_COLUMNS
                        self.assertEqual(list(book.active.values), [tuple(columns)])
                    finally:
                        book.close()

    def test_blank_codes_report_original_row(self):
        for module in [out_screen, ud_screen]:
            for blank in ["", "  ", "\t", "\u3000"]:
                with self.subTest(module=module.__name__, blank=blank):
                    path = self.csv(module, [self.row("1000001"), self.row(code=blank)])
                    with self.assertRaisesRegex(ValueError, "case.csv：第 3 .*drug_id 空白"):
                        module.read_screened_csv(path)

    def test_illegal_control_characters(self):
        characters = [chr(i) for i in [*range(9), 11, 12, *range(14, 32)]]
        for module in [out_screen, ud_screen]:
            for column in ["drug_id", "drug_name"]:
                for character in characters:
                    with self.subTest(module=module.__name__, column=column, character=ord(character)):
                        row = self.row()
                        row[column] = "bad" + character + "text"
                        with self.assertRaisesRegex(ValueError, column + " 含 Excel 不允許的控制字元"):
                            module.read_screened_csv(self.csv(module, [row]))

    def test_legal_text_and_formulas(self):
        for module in [out_screen, ud_screen]:
            with self.subTest(module=module.__name__):
                self.csv(module, [self.row(code="=001", name="=藥品\t名稱\n第二行")], name="11501.csv")
                with contextlib.redirect_stdout(io.StringIO()):
                    module.screen_to_excel(self.base, self.base / "output.xlsx")
                book = load_workbook(self.base / "output.xlsx")
                try:
                    self.assertEqual(book.active["A2"].data_type, "s")
                    self.assertEqual(book.active["B2"].value, "=藥品\t名稱\n第二行")
                    self.assertEqual(book.active["B2"].data_type, "s")
                finally:
                    book.close()

    def test_failure_keeps_previous_output(self):
        for module in [out_screen, ud_screen]:
            for bad in [self.row(code=""), self.row(name="bad\x01"), self.row("NaN")]:
                with self.subTest(module=module.__name__, bad=bad):
                    self.csv(module, [self.row()], name="11501.csv")
                    self.csv(module, [bad], name="11502.csv")
                    output = self.base / "output.xlsx"
                    output.write_bytes(b"previous-output")
                    with contextlib.redirect_stdout(io.StringIO()), self.assertRaises(ValueError):
                        module.screen_to_excel(self.base, output)
                    self.assertEqual(output.read_bytes(), b"previous-output")

    def test_ud_inventory_period_as_excel_text(self):
        for name in ["11501", "11509", "11512", "1151", "00102"]:
            self.csv(ud_screen, [self.row(), self.row(code="002")], name=f"{name}.csv")
        output = self.base / "output.xlsx"
        with contextlib.redirect_stdout(io.StringIO()):
            counts = ud_screen.screen_to_excel(self.base, output)
        self.assertEqual(sum(counts.values()), 10)
        book = load_workbook(output, read_only=True)
        try:
            self.assertEqual(book.sheetnames, ["00102", "11501", "11509", "1151", "11512"])
            for sheet in book:
                rows = list(sheet.rows)
                self.assertEqual([cell.value for cell in rows[0]], ud_screen.OUTPUT_COLUMNS)
                for row in rows[1:]:
                    self.assertEqual(row[5].value, sheet.title[:3])
                    self.assertEqual(row[6].value, sheet.title[3:].zfill(2))
                    self.assertEqual([cell.data_type for cell in row[5:]], ["s", "s"])
                    self.assertIsInstance(row[5].value, str)
                    self.assertIsInstance(row[6].value, str)
        finally:
            book.close()

    def test_ud_invalid_sheet_period_keeps_output(self):
        for index, name in enumerate(["case", "11500", "11513", "115", "115001", "11a01", "11501_2"]):
            with self.subTest(name=name):
                directory = self.base / str(index)
                directory.mkdir()
                pd.DataFrame([self.row()]).to_csv(directory / f"{name}.csv", index=False)
                output = directory / "output.xlsx"
                output.write_bytes(b"previous-output")
                with self.assertRaisesRegex(ValueError, "工作表"):
                    ud_screen.screen_to_excel(directory, output)
                self.assertEqual(output.read_bytes(), b"previous-output")

    def test_out_inventory_period_as_excel_text(self):
        for name in ["11501", "11509", "11512", "1151", "00102"]:
            self.csv(out_screen, [self.row("1000001"), self.row("-1.25"),
                                  self.row(code="002")], name=f"{name}.csv")
        output = self.base / "output.xlsx"
        with contextlib.redirect_stdout(io.StringIO()):
            counts = out_screen.screen_to_excel(self.base, output)
        self.assertEqual(sum(counts.values()), 10)
        book = load_workbook(output, read_only=True)
        try:
            self.assertEqual(book.sheetnames, ["00102", "11501", "11509", "1151", "11512"])
            for sheet in book:
                rows = list(sheet.rows)
                self.assertEqual([cell.value for cell in rows[0]], out_screen.OUTPUT_COLUMNS)
                self.assertEqual([row[2].value for row in rows[1:]], [-1.25, 1])
                for row in rows[1:]:
                    self.assertEqual(row[4].value, sheet.title[:3])
                    self.assertEqual(row[5].value, sheet.title[3:].zfill(2))
                    self.assertEqual([cell.data_type for cell in row[4:]], ["s", "s"])
                    self.assertIsInstance(row[4].value, str)
                    self.assertIsInstance(row[5].value, str)
        finally:
            book.close()

    def test_out_invalid_sheet_period_keeps_output(self):
        for index, name in enumerate(["case", "11500", "11513", "115", "115001", "11a01", "11501_2"]):
            with self.subTest(name=name):
                directory = self.base / str(index)
                directory.mkdir()
                pd.DataFrame([self.row()]).to_csv(directory / "00101.csv", index=False)
                pd.DataFrame([self.row()]).to_csv(directory / f"{name}.csv", index=False)
                output = directory / "output.xlsx"
                output.write_bytes(b"previous-output")
                with contextlib.redirect_stdout(io.StringIO()), self.assertRaisesRegex(ValueError, "工作表"):
                    out_screen.screen_to_excel(directory, output)
                self.assertEqual(output.read_bytes(), b"previous-output")

    def books(self, rows):
        for name in total_screen.INPUT_NAMES:
            pd.DataFrame(rows, columns=total_screen.COLUMNS).to_excel(
                self.base / name, sheet_name="11501", index=False
            )

    def test_total_periods_remain_separate_and_are_excel_text(self):
        for name in total_screen.INPUT_NAMES:
            with pd.ExcelWriter(self.base / name, engine="openpyxl") as writer:
                for sheet, qty in [("00102", 1.25), ("1151", 3), ("11501", -1), ("11512", 10)]:
                    pd.DataFrame([("001", qty)], columns=total_screen.COLUMNS).to_excel(
                        writer, sheet_name=sheet, index=False
                    )
        with contextlib.redirect_stdout(io.StringIO()):
            result = total_screen.total_to_excel(self.base)
        self.assertEqual(result.columns.tolist(), total_screen.OUTPUT_COLUMNS)
        self.assertEqual(result.inv_year.tolist(), ["001", "115", "115"])
        self.assertEqual(result.inv_month.tolist(), ["02", "01", "12"])
        self.assertEqual(result.total_qty.tolist(), [2.5, 4, 20])
        for column in ("inv_year", "inv_month"):
            self.assertEqual(str(result[column].dtype), "string")
        book = load_workbook(self.base / total_screen.OUTPUT_NAME)
        try:
            for row in book.active.iter_rows(min_row=2):
                self.assertEqual(row[2].data_type, "s")
                self.assertEqual(row[3].data_type, "s")
                self.assertEqual(len(row[2].value), 3)
                self.assertEqual(len(row[3].value), 2)
            self.assertEqual(book.active["C2"].value, "001")
            self.assertEqual(book.active["D2"].value, "02")
        finally:
            book.close()

    def test_total_invalid_period_keeps_output(self):
        for sheet in ("Sheet1", "11500", "11513", "115001", "11A01"):
            with self.subTest(sheet=sheet):
                for name in total_screen.INPUT_NAMES:
                    pd.DataFrame(columns=total_screen.COLUMNS).to_excel(
                        self.base / name, sheet_name=sheet, index=False
                    )
                output = self.base / total_screen.OUTPUT_NAME
                output.write_bytes(b"previous-output")
                with self.assertRaisesRegex(ValueError, "工作表"):
                    total_screen.total_to_excel(self.base)
                self.assertEqual(output.read_bytes(), b"previous-output")

    def test_total_filters_input_but_not_sum(self):
        self.books([("001", "1000000"), ("001", "1000000"),
                    ("", "5000000000000000000"), ("001", "-9223372036854775808")])
        with contextlib.redirect_stdout(io.StringIO()) as log:
            result = total_screen.total_to_excel(self.base)
        self.assertEqual(result.drug_id.tolist(), ["001"])
        self.assertEqual(result.total_qty.tolist(), [4000000])
        self.assertEqual(log.getvalue().count("數量超出範圍忽略 2 筆"), 2)

    def test_total_all_ignored_is_valid_empty_workbook(self):
        self.books([("001", "1e400")])
        with contextlib.redirect_stdout(io.StringIO()):
            result = total_screen.total_to_excel(self.base)
        self.assertTrue(result.empty)
        self.assertTrue(pd.read_excel(self.base / total_screen.OUTPUT_NAME).empty)

    def test_total_invalid_input_keeps_output(self):
        for rows in [[("", "1")], [("001", "NaN")]]:
            with self.subTest(rows=rows):
                self.books(rows)
                output = self.base / total_screen.OUTPUT_NAME
                output.write_bytes(b"previous-output")
                with self.assertRaises(ValueError):
                    total_screen.total_to_excel(self.base)
                self.assertEqual(output.read_bytes(), b"previous-output")


if __name__ == "__main__":
    unittest.main()
