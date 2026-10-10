"""從輸入活頁簿到輸出活頁簿驗證月明細合併與失敗保護。"""

import contextlib
import io
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

import summay


class SummayTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        fixed_directory = patch.object(summay, "BASE_DIR", self.root)
        fixed_directory.start()
        self.addCleanup(fixed_directory.stop)
        self.base = self.root / "base.xlsx"
        self.monthly = self.root / "月報表.xlsx"
        self.total = self.root / "total_screen.xlsx"
        self.output = self.root / "summay.xlsx"
        self.base_rows = [
            ["ABC01O", "Alpha", "O"], ["ABC02I", "Beta", "I"],
            ["ABC03E", "Gamma", "E"], ["ABC04O", "=literal name", "O"],
            ["ABC05S", "Epsilon", 7],
        ]
        self.month_rows = {
            "1161": [["ABC02I", 0, 8], ["ABC01O", 10, 3], ["ABC03E", 4, -2],
                     ["ABC04O", -2, 4], ["ABC05S", 6, 10]],
            "11512": [["ABC01O", 2, 5], ["ABC04O", 1, 0]],
            "11602": [],
        }
        self.total_rows = [
            ["ABC01O", 7, "116", "01"], ["ABC02I", 0, "116", "01"],
            ["ABC03E", -3, "116", "01"], ["ABC04O", 5, "116", "01"],
            ["ABC01O", 3, "115", "12"], ["ABC04O", 1, "115", "12"],
            ["ABC03E", 12, "115", "12"],
        ]
        self.write_inputs()

    def workbook(self, path, columns, sheets):
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            for sheet, rows in sheets.items():
                pd.DataFrame(rows, columns=columns).to_excel(writer, sheet_name=sheet, index=False)
        # 測試來源的 = 字串必須真的存為文字，而不是 Excel 公式。
        book = load_workbook(path)
        for sheet in book:
            for row in sheet:
                for cell in row:
                    if cell.data_type == "f":
                        cell.data_type = "s"
        book.save(path)
        book.close()

    def write_inputs(self):
        self.workbook(self.base, summay.BASE_COLUMNS, {"base": self.base_rows})
        self.workbook(self.monthly, summay.MONTHLY_COLUMNS, self.month_rows)
        self.workbook(self.total, summay.TOTAL_COLUMNS, {"total_screen": self.total_rows})

    def run_merge(self):
        with contextlib.redirect_stdout(io.StringIO()):
            return summay.merge_to_excel()

    def test_merge_periods_filtering_and_chinese_output(self):
        counts = self.run_merge()
        self.assertEqual(counts["1161"], {"input": 5, "excluded": 2, "filled_zero": 1, "output": 3})
        book = load_workbook(self.output)
        try:
            self.assertEqual(book.sheetnames, list(self.month_rows))
            self.assertEqual(list(book["1161"].values), [
                tuple(summay.OUTPUT_NAMES.values()),
                ("ABC01O", "Alpha", 10, 7, 3, "116", "01", "O"),
                ("ABC03E", "Gamma", 4, -3, -2, "116", "01", "E"),
                ("ABC05S", "Epsilon", 6, 0, 10, "116", "01", "7"),
            ])
            self.assertEqual(list(book["11512"].values)[1:], [
                ("ABC01O", "Alpha", 2, 3, 5, "115", "12", "O"),
                ("ABC04O", "=literal name", 1, 1, 0, "115", "12", "O"),
            ])
            self.assertEqual(book["11512"]["B3"].data_type, "s")
            self.assertEqual(book["1161"]["G2"].data_type, "s")
            self.assertEqual(book["1161"]["D2"].data_type, "n")
            self.assertEqual(list(book["11602"].values), [tuple(summay.OUTPUT_NAMES.values())])
        finally:
            book.close()

    def test_duplicate_totals_report_all_rows_across_sheets_and_keep_old_output(self):
        self.output.write_bytes(b"old output")
        repeated = [["ABC01O", i, "116", "1"] for i in range(12)]
        self.workbook(self.total, summay.TOTAL_COLUMNS,
                      {"total_screen": self.total_rows, "more": repeated})
        with self.assertRaises(ValueError) as error:
            self.run_merge()
        message = str(error.exception)
        self.assertIn("1 組、13 列", message)
        self.assertEqual(message.count("ABC01O"), 13)
        self.assertEqual(message.count("more"), 12)
        self.assertIn("來源工作表", message)
        self.assertIn("Excel 列號", message)
        self.assertEqual(self.output.read_bytes(), b"old output")
        self.assertEqual(list(self.root.glob(".summay-*.xlsx")), [])

    def test_preprocessed_sources_still_reject_duplicate_keys(self):
        for source in ("base", "monthly"):
            with self.subTest(source=source):
                self.write_inputs()
                self.output.write_bytes(b"old output")
                if source == "base":
                    self.workbook(self.base, summay.BASE_COLUMNS,
                                  {"base": self.base_rows + [self.base_rows[0]]})
                else:
                    self.workbook(self.monthly, summay.MONTHLY_COLUMNS,
                                  {"1161": self.month_rows["1161"] * 2})
                with self.assertRaisesRegex(ValueError, "重複鍵"):
                    self.run_merge()
                self.assertEqual(self.output.read_bytes(), b"old output")

    def test_invalid_inputs_never_replace_old_output(self):
        for column, bad_value in (("total_qty", "not a number"), ("inv_month", "13")):
            with self.subTest(column=column):
                self.write_inputs()
                rows = [row[:] for row in self.total_rows]
                rows[0][summay.TOTAL_COLUMNS.index(column)] = bad_value
                self.workbook(self.total, summay.TOTAL_COLUMNS, {"total_screen": rows})
                self.output.write_bytes(b"old output")
                with self.assertRaisesRegex(ValueError, column):
                    self.run_merge()
                self.assertEqual(self.output.read_bytes(), b"old output")

    def test_nonpositive_months_retain_headers_and_missing_usage_fills_zero(self):
        self.workbook(self.monthly, summay.MONTHLY_COLUMNS,
                      {"11501": [["ABC01O", -1, 5]], "11502": [["ABC01O", 0, -2]], "11503": [["ABC01O", 1, -2]]})
        self.workbook(self.total, summay.TOTAL_COLUMNS, {"total_screen": []})
        counts = self.run_merge()
        self.assertEqual(counts["11501"]["output"], 0)
        self.assertEqual(counts["11502"]["output"], 0)
        self.assertEqual(counts["11503"]["filled_zero"], 1)
        book = load_workbook(self.output, read_only=True)
        try:
            self.assertEqual(list(book["11501"].values), [tuple(summay.OUTPUT_NAMES.values())])
            self.assertEqual(list(book["11502"].values), [tuple(summay.OUTPUT_NAMES.values())])
            self.assertEqual(list(book["11503"].values)[1],
                             ("ABC01O", "Alpha", 1, 0, -2, "115", "03", "O"))
        finally:
            book.close()

    def test_failed_write_preserves_old_output_and_removes_temporary_file(self):
        self.output.write_bytes(b"old output")
        with patch.object(summay, "format_sheet", side_effect=OSError("write failed")):
            with self.assertRaisesRegex(OSError, "write failed"):
                self.run_merge()
        self.assertEqual(self.output.read_bytes(), b"old output")
        self.assertEqual(list(self.root.glob(".summay-*.xlsx")), [])

    def test_one_click_execution_uses_fixed_directory_and_protects_input_files(self):
        previous = Path.cwd()
        other_directory = self.root / "other"
        other_directory.mkdir()
        try:
            os.chdir(other_directory)
            with contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(summay.main(), 0)
        finally:
            os.chdir(previous)
        self.assertTrue(self.output.is_file())
        self.assertEqual(list(other_directory.iterdir()), [])
        before = self.monthly.read_bytes()
        errors = io.StringIO()
        with patch.object(summay, "OUTPUT_NAME", "月報表.xlsx"), contextlib.redirect_stderr(errors):
            self.assertEqual(summay.main(), 1)
        self.assertIn("不可與任何輸入檔案相同", errors.getvalue())
        self.assertEqual(self.monthly.read_bytes(), before)


if __name__ == "__main__":
    unittest.main()
