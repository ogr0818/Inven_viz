#!/usr/bin/env python3
"""以 SQLAlchemy 管理藥品主檔及民國年月資料；Excel 僅讀取、不修改。"""

from __future__ import annotations

import argparse
from collections import Counter
from contextlib import closing
from dataclasses import asdict, dataclass
from datetime import datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
import json
import os
import re
from pathlib import Path
import sqlite3
import sys
import tempfile
from typing import Any
from zipfile import BadZipFile

try:
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException
    from sqlalchemy import (
        CheckConstraint, ForeignKey, Index, Integer, Text, URL,
        create_engine, delete, event, func, inspect, insert, select,
    )
    from sqlalchemy.dialects.sqlite import insert as sqlite_insert
    from sqlalchemy.engine import Connection, Engine
    from sqlalchemy.exc import SQLAlchemyError
    from sqlalchemy.orm import DeclarativeBase, Mapped, mapped_column
    from sqlalchemy.schema import CreateIndex, CreateTable
except ImportError as exc:
    raise SystemExit(
        "缺少必要套件，請先在專案目錄執行：uv sync\n"
        f"詳細原因：{exc}"
    ) from exc


SCHEMA_VERSION = 1
ROOT = Path(__file__).resolve().parent
BASE_HEADERS = ("drug_id", "drug_name", "drug_type")
MONTH_HEADERS = (
    "藥品代碼", "藥品名稱", "實發量", "住院耗用", "庫存量",
    "撥補年份", "撥補月份", "劑型",
)


class InputError(ValueError):
    """可由使用者修正的來源資料或資料庫錯誤。"""


class Base(DeclarativeBase):
    pass


class DrugCode(Base):
    __tablename__ = "drug_codes"
    drug_code: Mapped[str] = mapped_column(Text, primary_key=True, nullable=False)
    __table_args__ = (
        CheckConstraint("length(trim(drug_code)) > 0", name="ck_drug_code_nonempty"),
    )


class DrugMaster(Base):
    __tablename__ = "drug_master"
    drug_code: Mapped[str] = mapped_column(
        Text, ForeignKey("drug_codes.drug_code", ondelete="RESTRICT"),
        primary_key=True, nullable=False,
    )
    drug_name: Mapped[str] = mapped_column(Text, nullable=False)
    drug_type: Mapped[str] = mapped_column(Text, nullable=False)


class MonthlySummary(Base):
    __tablename__ = "monthly_summary"
    drug_code: Mapped[str] = mapped_column(
        Text, ForeignKey("drug_codes.drug_code", ondelete="RESTRICT"),
        primary_key=True, nullable=False,
    )
    roc_year: Mapped[int] = mapped_column(Integer, primary_key=True, nullable=False)
    roc_month: Mapped[int] = mapped_column(Integer, primary_key=True, nullable=False)
    drug_name: Mapped[str] = mapped_column(Text, nullable=False)
    drug_type: Mapped[str] = mapped_column(Text, nullable=False)
    issued_quantity: Mapped[int] = mapped_column(Integer, nullable=False)
    inpatient_usage: Mapped[int] = mapped_column(Integer, nullable=False)
    stock_quantity: Mapped[int] = mapped_column(Integer, nullable=False)
    __table_args__ = (
        CheckConstraint("roc_year > 0", name="ck_monthly_roc_year"),
        CheckConstraint("roc_month BETWEEN 1 AND 12", name="ck_monthly_roc_month"),
        CheckConstraint("issued_quantity > 0", name="ck_monthly_issued_positive"),
        Index("ix_monthly_summary_period", "roc_year", "roc_month"),
    )


@dataclass(frozen=True)
class MasterRow:
    drug_code: str
    drug_name: str
    drug_type: str


@dataclass
class MonthlyImport:
    rows: list[dict[str, Any]]
    periods: set[tuple[int, int]]
    all_codes: set[str]
    locations: dict[str, str]
    source_rows: int
    sheets: list[str]

    def report(self) -> dict[str, Any]:
        counts = Counter((r["roc_year"], r["roc_month"]) for r in self.rows)
        return {
            "來源筆數": self.source_rows,
            "匯入筆數": len(self.rows),
            "排除筆數": self.source_rows - len(self.rows),
            "工作表": self.sheets,
            "各月匯入筆數": {
                f"{year}/{month:02d}": counts[(year, month)]
                for year, month in sorted(self.periods)
            },
            "負數住院耗用筆數": sum(r["inpatient_usage"] < 0 for r in self.rows),
            "負數庫存筆數": sum(r["stock_quantity"] < 0 for r in self.rows),
        }


def text_value(value: Any, label: str, location: str) -> str:
    if value is None or isinstance(value, bool):
        raise InputError(f"{location}：{label}不可空白，且不得為布林值。")
    # 數字分類代碼 0–9 也以文字保存，不自行解讀代碼含義。
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    result = str(value).strip()
    if not result:
        raise InputError(f"{location}：{label}不可空白。")
    return result


def integer_value(value: Any, label: str, location: str) -> int:
    if value is None or isinstance(value, bool):
        raise InputError(f"{location}：{label}必須為整數。")
    try:
        number = Decimal(str(value).strip())
    except (InvalidOperation, ValueError) as exc:
        raise InputError(f"{location}：{label}必須為整數。") from exc
    if not number.is_finite() or number != number.to_integral_value():
        raise InputError(f"{location}：{label}必須為整數，不會自動四捨五入。")
    result = int(number)
    if not -(2**63) <= result < 2**63:
        raise InputError(f"{location}：{label}超出 SQLite 整數範圍。")
    return result


def workbook_rows(path: Path, headers: tuple[str, ...], empty_sheets: list[str] | None = None):
    """逐張處理資料表；允許欄位重排及額外欄位，拒絕公式與重複欄名。"""
    if not path.is_file():
        raise InputError(f"找不到 Excel 檔案：{path}")
    if path.suffix.lower() != ".xlsx":
        raise InputError(f"只支援 .xlsx 檔案：{path}")
    try:
        book = load_workbook(path, read_only=True, data_only=False)
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError) as exc:
        raise InputError(f"無法讀取 Excel：{path.name}；{exc}") from exc
    try:
        for sheet in book.worksheets:
            iterator = sheet.iter_rows()
            header_cells = next(iterator, ())
            if not any(c.value is not None for c in header_cells):
                if any(any(c.value is not None for c in row) for row in iterator):
                    raise InputError(f"{path.name} / {sheet.title}：第一列必須為欄位名稱。")
                continue
            names = [str(c.value).strip() if c.value is not None else "" for c in header_cells]
            missing = [name for name in headers if name not in names]
            duplicates = [name for name in headers if names.count(name) > 1]
            if missing or duplicates:
                raise InputError(
                    f"{path.name} / {sheet.title}："
                    + (f"缺少欄位 {', '.join(missing)}。" if missing else "")
                    + (f"重複欄位 {', '.join(duplicates)}。" if duplicates else "")
                )
            indices = [names.index(name) for name in headers]
            has_rows = False
            for row_number, cells in enumerate(iterator, start=2):
                if not any(c.value is not None for c in cells):
                    continue
                has_rows = True
                location = f"{path.name} / {sheet.title} / 第 {row_number} 列"
                chosen = [cells[i] if i < len(cells) else None for i in indices]
                for label, cell in zip(headers, chosen):
                    if cell is not None and cell.data_type in ("f", "e"):
                        raise InputError(f"{location}：{label}含公式或 Excel 錯誤，請先提供清理後的數值。")
                yield location, sheet.title, [c.value if c is not None else None for c in chosen]
            if not has_rows and empty_sheets is not None:
                empty_sheets.append(sheet.title)
    finally:
        book.close()


def read_master(path: Path) -> list[MasterRow]:
    rows: list[MasterRow] = []
    seen: dict[str, str] = {}
    for location, _, values in workbook_rows(path, BASE_HEADERS):
        code, name, kind = [text_value(v, h, location) for v, h in zip(values, BASE_HEADERS)]
        if code in seen:
            raise InputError(f"{location}：重複藥碼 {code}；首次出現在 {seen[code]}。")
        seen[code] = location
        rows.append(MasterRow(code, name, kind))
    if not rows:
        raise InputError("主檔沒有藥品資料，拒絕以空主檔替換資料庫。")
    return rows


def read_monthly(path: Path) -> MonthlyImport:
    result = MonthlyImport([], set(), set(), {}, 0, [])
    seen: dict[tuple[str, int, int], str] = {}
    empty_sheets: list[str] = []
    for location, sheet, values in workbook_rows(path, MONTH_HEADERS, empty_sheets):
        code, name, issued, usage, stock, year, month, kind = values
        code = text_value(code, "藥品代碼", location)
        name = text_value(name, "藥品名稱", location)
        kind = text_value(kind, "劑型", location)
        issued = integer_value(issued, "實發量", location)
        usage = integer_value(usage, "住院耗用", location)
        stock = integer_value(stock, "庫存量", location)
        year = integer_value(year, "撥補年份", location)
        month = integer_value(month, "撥補月份", location)
        if year <= 0 or not 1 <= month <= 12:
            raise InputError(f"{location}：民國年須大於零，月份須介於 1 與 12。")
        key = (code, year, month)
        if key in seen:
            raise InputError(f"{location}：重複藥碼年月 {key}；首次出現在 {seen[key]}。")
        seen[key] = location
        result.source_rows += 1
        # 必須在實發量篩選前登記月份；整月皆為零時也需清除原有月份。
        result.periods.add((year, month))
        result.all_codes.add(code)
        result.locations.setdefault(code, location)
        if sheet not in result.sheets:
            result.sheets.append(sheet)
        if issued > 0:
            result.rows.append({
                "drug_code": code, "roc_year": year, "roc_month": month,
                "drug_name": name, "drug_type": kind, "issued_quantity": issued,
                "inpatient_usage": usage, "stock_quantity": stock,
            })
    # 正實發量篩選後的整月空表仍代表整月替換，避免舊月紀錄殘留。
    for sheet in empty_sheets:
        if not re.fullmatch(r"[0-9]{4,5}", sheet):
            raise InputError(f"{path.name} / {sheet}：空月表須以三位民國年及一或兩位月份命名。")
        year, month = int(sheet[:3]), int(sheet[3:])
        if year <= 0 or not 1 <= month <= 12:
            raise InputError(f"{path.name} / {sheet}：空月表年月不合法。")
        result.periods.add((year, month))
        result.sheets.append(sheet)
    if not result.periods:
        raise InputError("月檔沒有可辨識年月的資料列；無法決定要替換哪個月份。")
    return result


def require_known_codes(data: MonthlyImport, master_codes: set[str]) -> None:
    unknown = sorted(data.all_codes - master_codes)
    if unknown:
        sample = "；".join(f"{code}（{data.locations[code]}）" for code in unknown[:10])
        raise InputError(f"月檔有 {len(unknown)} 個藥碼不在目前主檔：{sample}。請先清理資料或更新主檔。")


def database_engine(path: Path) -> Engine:
    engine = create_engine(URL.create("sqlite", database=str(path)), connect_args={"timeout": 30})

    @event.listens_for(engine, "connect")
    def configure_sqlite(dbapi_connection, _record):
        # SQLAlchemy 統一控制 BEGIN，使 DDL、刪除及新增可真正一起回復。
        dbapi_connection.isolation_level = None
        cursor = dbapi_connection.cursor()
        cursor.execute("PRAGMA foreign_keys = ON")
        cursor.execute("PRAGMA busy_timeout = 30000")
        cursor.close()

    @event.listens_for(engine, "begin")
    def begin_transaction(connection):
        statement = "BEGIN IMMEDIATE" if connection.get_execution_options().get("write_transaction") else "BEGIN"
        connection.exec_driver_sql(statement)

    return engine


def require_database(path: Path) -> None:
    if not path.is_file():
        raise InputError(f"資料庫尚不存在：{path}；請先使用 init 建立。")


def validate_schema(connection: Connection) -> None:
    version = connection.exec_driver_sql("PRAGMA user_version").scalar_one()
    if version != SCHEMA_VERSION:
        raise InputError(f"資料庫 schema 版本為 {version}，程式只支援 {SCHEMA_VERSION}；不會自動覆寫。")
    inspector = inspect(connection)
    expected = set(Base.metadata.tables)
    if set(inspector.get_table_names()) != expected:
        raise InputError("資料庫資料表與預期 schema 不一致，不會進行更新。")
    for table in Base.metadata.sorted_tables:
        columns = {c["name"]: c for c in inspector.get_columns(table.name)}
        if set(columns) != set(table.columns.keys()):
            raise InputError(f"資料表 {table.name} 的欄位與預期 schema 不一致。")
        for column in table.columns:
            actual = columns[column.name]
            expected_type = Integer if isinstance(column.type, Integer) else Text
            if actual["nullable"] or not isinstance(actual["type"], expected_type):
                raise InputError(f"{table.name}.{column.name} 的型別或非空限制不符合 schema。")
        primary_key = inspector.get_pk_constraint(table.name)["constrained_columns"]
        if primary_key != [c.name for c in table.primary_key.columns]:
            raise InputError(f"資料表 {table.name} 的主鍵與預期 schema 不一致。")
        expected_checks = {str(c.sqltext) for c in table.constraints if isinstance(c, CheckConstraint)}
        actual_checks = {c["sqltext"] for c in inspector.get_check_constraints(table.name)}
        if expected_checks != actual_checks:
            raise InputError(f"資料表 {table.name} 的檢核限制與預期 schema 不一致。")
        if table.name != "drug_codes":
            fks = inspector.get_foreign_keys(table.name)
            if len(fks) != 1 or fks[0]["constrained_columns"] != ["drug_code"] or fks[0]["referred_table"] != "drug_codes" or fks[0]["referred_columns"] != ["drug_code"] or fks[0]["options"].get("ondelete") != "RESTRICT":
                raise InputError(f"資料表 {table.name} 的外鍵與預期 schema 不一致。")


def validate_integrity(connection: Connection) -> None:
    checks = connection.exec_driver_sql("PRAGMA quick_check").scalars().all()
    if checks != ["ok"]:
        raise InputError(f"SQLite 完整性檢查失敗：{checks}")
    if connection.exec_driver_sql("PRAGMA foreign_key_check").first() is not None:
        raise InputError("資料庫存在外鍵失去對應的紀錄。")
    if connection.exec_driver_sql("PRAGMA foreign_keys").scalar_one() != 1:
        raise InputError("SQLite 外鍵檢查未啟用。")


def backup_database(path: Path) -> Path:
    folder = path.parent / "backups"
    folder.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(timezone(timedelta(hours=8))).strftime("%Y%m%d-%H%M%S-%f")
    target = folder / f"{path.stem}.{stamp}.db"
    try:
        with closing(sqlite3.connect(path.as_uri() + "?mode=ro", uri=True, timeout=30)) as source:
            with closing(sqlite3.connect(target)) as destination:
                source.backup(destination)
    except Exception:
        target.unlink(missing_ok=True)
        raise
    return target


def replace_master_rows(connection: Connection, rows: list[MasterRow]) -> None:
    connection.execute(
        sqlite_insert(DrugCode).on_conflict_do_nothing(index_elements=["drug_code"]),
        [{"drug_code": r.drug_code} for r in rows],
    )
    connection.execute(delete(DrugMaster))
    connection.execute(insert(DrugMaster), [asdict(r) for r in rows])


def replace_month_rows(connection: Connection, data: MonthlyImport) -> None:
    for year, month in sorted(data.periods):
        connection.execute(delete(MonthlySummary).where(
            MonthlySummary.roc_year == year, MonthlySummary.roc_month == month,
        ))
    if data.rows:
        connection.execute(insert(MonthlySummary), data.rows)


def database_info(connection: Connection) -> dict[str, Any]:
    counts = {
        model.__tablename__: connection.scalar(select(func.count()).select_from(model))
        for model in (DrugCode, DrugMaster, MonthlySummary)
    }
    periods = connection.execute(select(
        MonthlySummary.roc_year, MonthlySummary.roc_month, func.count(),
    ).group_by(MonthlySummary.roc_year, MonthlySummary.roc_month)
      .order_by(MonthlySummary.roc_year, MonthlySummary.roc_month)).all()
    missing_master = connection.scalar(select(func.count()).select_from(MonthlySummary)
        .outerjoin(DrugMaster, MonthlySummary.drug_code == DrugMaster.drug_code)
        .where(DrugMaster.drug_code.is_(None)))
    return {
        "schema版本": SCHEMA_VERSION, "資料表筆數": counts,
        "各月筆數": {f"{y}/{m:02d}": count for y, m, count in periods},
        "目前主檔未收錄的歷史月紀錄": missing_master,
    }


def initialize(db: Path, base: Path, summary: Path, dry_run: bool = False) -> dict[str, Any]:
    if db.exists():
        raise InputError(f"資料庫已存在：{db}；init 不會覆寫，請改用更新指令。")
    masters = read_master(base)
    monthly = read_monthly(summary)
    require_known_codes(monthly, {r.drug_code for r in masters})
    return initialize_data(db, masters, monthly, dry_run)


def initialize_data(db: Path, masters: list[MasterRow], monthly: MonthlyImport,
                    dry_run: bool = False) -> dict[str, Any]:
    """以已驗證月明細建立 DB；保留拒絕覆寫及原子發布保護。"""
    if db.exists():
        raise InputError(f"資料庫已存在：{db}；init 不會覆寫，請改用更新指令。")
    require_known_codes(monthly, {r.drug_code for r in masters})
    report = {"作業": "建立資料庫", "資料庫": str(db), "預演": dry_run,
              "主檔筆數": len(masters), **monthly.report()}
    if dry_run:
        return report
    db.parent.mkdir(parents=True, exist_ok=True)
    handle, temp_name = tempfile.mkstemp(prefix=f".{db.name}.init-", suffix=".tmp", dir=db.parent)
    os.close(handle)
    temp_path = Path(temp_name)
    engine = database_engine(temp_path)
    try:
        with engine.connect().execution_options(write_transaction=True) as connection:
            with connection.begin():
                Base.metadata.create_all(connection)
                connection.exec_driver_sql(f"PRAGMA user_version = {SCHEMA_VERSION}")
                replace_master_rows(connection, masters)
                replace_month_rows(connection, monthly)
                validate_integrity(connection)
                report.update(database_info(connection))
        engine.dispose()
        # 同資料夾硬連結以原子方式發布；既有路徑會拒絕覆寫。
        os.link(temp_path, db)
    finally:
        engine.dispose()
        temp_path.unlink(missing_ok=True)
        Path(str(temp_path) + "-journal").unlink(missing_ok=True)
    return report


def update_monthly(db: Path, summary: Path, dry_run: bool = False) -> dict[str, Any]:
    require_database(db)
    monthly = read_monthly(summary)
    return update_monthly_data(db, monthly, dry_run)


def update_monthly_data(db: Path, monthly: MonthlyImport,
                        dry_run: bool = False) -> dict[str, Any]:
    """共用整月交易，接受 Excel 讀入或三檔整併後的月明細。"""
    require_database(db)
    report = {"作業": "整月替換", "資料庫": str(db), "預演": dry_run, **monthly.report()}
    engine = database_engine(db)
    try:
        with engine.connect().execution_options(write_transaction=not dry_run) as connection:
            with connection.begin():
                validate_schema(connection)
                validate_integrity(connection)
                codes = set(connection.scalars(select(DrugMaster.drug_code)))
                require_known_codes(monthly, codes)
                old_counts = {
                    f"{y}/{m:02d}": connection.scalar(select(func.count()).select_from(MonthlySummary)
                        .where(MonthlySummary.roc_year == y, MonthlySummary.roc_month == m))
                    for y, m in sorted(monthly.periods)
                }
                report["各月替換前筆數"] = old_counts
                if not dry_run:
                    report["更新前備份"] = str(backup_database(db))
                    replace_month_rows(connection, monthly)
                    validate_integrity(connection)
                    report.update(database_info(connection))
    finally:
        engine.dispose()
    return report


def replace_base(db: Path, base: Path, dry_run: bool = False) -> dict[str, Any]:
    require_database(db)
    masters = read_master(base)
    engine = database_engine(db)
    report = {"作業": "完整替換主檔", "資料庫": str(db), "預演": dry_run, "新主檔筆數": len(masters)}
    try:
        with engine.connect().execution_options(write_transaction=not dry_run) as connection:
            with connection.begin():
                validate_schema(connection)
                validate_integrity(connection)
                old_codes = set(connection.scalars(select(DrugMaster.drug_code)))
                new_codes = {r.drug_code for r in masters}
                report.update({"新增主檔藥碼數": len(new_codes - old_codes),
                               "移出主檔藥碼數": len(old_codes - new_codes)})
                if not dry_run:
                    report["更新前備份"] = str(backup_database(db))
                    replace_master_rows(connection, masters)
                    validate_integrity(connection)
                    report.update(database_info(connection))
    finally:
        engine.dispose()
    return report


def information(db: Path) -> dict[str, Any]:
    require_database(db)
    engine = database_engine(db)
    try:
        with engine.connect() as connection:
            validate_schema(connection)
            validate_integrity(connection)
            return {"資料庫": str(db), "完整性檢查": "通過", **database_info(connection)}
    finally:
        engine.dispose()


def schema_sql() -> str:
    engine = database_engine(Path(":memory:"))
    try:
        statements = ["-- schema 版本 1；由 manage_summary.py 的 SQLAlchemy 模型產生。",
                      "PRAGMA foreign_keys = ON;", f"PRAGMA user_version = {SCHEMA_VERSION};"]
        for table in Base.metadata.sorted_tables:
            statements.append(str(CreateTable(table).compile(dialect=engine.dialect)).strip() + ";")
            for index in sorted(table.indexes, key=lambda i: i.name):
                statements.append(str(CreateIndex(index).compile(dialect=engine.dialect)).strip() + ";")
        return "\n\n".join(statements) + "\n"
    finally:
        engine.dispose()


def cli() -> int:
    parser = argparse.ArgumentParser(description="建立及更新 summary.db；年月保留民國年制。")
    commands = parser.add_subparsers(dest="command", required=True)
    for name, help_text in (
        ("init", "首次建立資料庫，不覆寫既有檔案"),
        ("update-monthly", "整月替換單張或多張工作表的月資料"),
        ("replace-base", "完整替換主檔，保留歷史月資料"),
        ("info", "檢查 schema、完整性與資料筆數"),
    ):
        command = commands.add_parser(name, help=help_text)
        command.add_argument("--db", type=Path, default=ROOT / "summary.db", help="資料庫路徑")
        if name in ("init", "replace-base"):
            command.add_argument("--base", type=Path, default=ROOT / "base.xlsx", help="完整藥品主檔")
        if name in ("init", "update-monthly"):
            default_summary = ROOT / "summary.xlsx"
            if name == "init" and not default_summary.exists():
                default_summary = ROOT / "summay.xlsx"
            command.add_argument("--summary", type=Path, default=default_summary, help="完整月份的 Excel 檔")
        if name != "info":
            command.add_argument("--dry-run", action="store_true", help="只驗證及顯示預計變更，不寫入或備份")
    args = parser.parse_args()
    try:
        db = args.db.expanduser().resolve()
        if args.command == "init":
            report = initialize(db, args.base.expanduser().resolve(), args.summary.expanduser().resolve(), args.dry_run)
        elif args.command == "update-monthly":
            report = update_monthly(db, args.summary.expanduser().resolve(), args.dry_run)
        elif args.command == "replace-base":
            report = replace_base(db, args.base.expanduser().resolve(), args.dry_run)
        else:
            report = information(db)
    except (InputError, SQLAlchemyError, OSError, sqlite3.Error, BadZipFile, InvalidFileException) as exc:
        print(f"作業失敗：{exc}\n未提交資料庫更新；請修正原因後重新執行。", file=sys.stderr)
        return 2
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(cli())

'''
manage_summary.py update-monthly --summary summary.xlsx
'''